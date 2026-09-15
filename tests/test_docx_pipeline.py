from __future__ import annotations

import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.style import WD_STYLE_TYPE


ROOT_DIR = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = ROOT_DIR / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from docx_pipeline import render_docx_from_template
from app_service import export_round_output
from app_service import run_round_for_app
from skill_round_helper import run_skill_round


class DocxPipelineTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = Path(tempfile.mkdtemp(prefix="baibai-docx-test-"))

    def tearDown(self) -> None:
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def write_template_docx(self, path: Path) -> None:
        document = Document()
        style = document.styles.add_style("Custom Body", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["Normal"]

        first = document.add_paragraph(style="Custom Body")
        first_run = first.add_run("原始第一段。")
        first_run.bold = True

        document.add_paragraph("原始第二段。")

        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "表格内容保持不变"

        section = document.sections[0]
        section.header.paragraphs[0].text = "页眉保留"
        document.save(str(path))

    def test_render_docx_from_template_replaces_text_and_preserves_document_structure(self) -> None:
        template_path = self.temp_dir / "template.docx"
        output_text_path = self.temp_dir / "result.txt"
        output_docx_path = self.temp_dir / "result.docx"
        self.write_template_docx(template_path)
        output_text_path.write_text("改写后的第一段。\n\n改写后的第二段。", encoding="utf-8")

        render_docx_from_template(
            template_path,
            output_text_path,
            output_docx_path,
        )

        document = Document(str(output_docx_path))
        self.assertEqual(
            [paragraph.text for paragraph in document.paragraphs],
            ["改写后的第一段。", "改写后的第二段。"],
        )
        self.assertEqual(document.paragraphs[0].style.name, "Custom Body")
        self.assertTrue(document.paragraphs[0].runs[0].bold)
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(document.tables[0].cell(0, 0).text, "表格内容保持不变")
        self.assertEqual(document.sections[0].header.paragraphs[0].text, "页眉保留")

    def test_render_docx_from_template_rejects_paragraph_count_mismatch(self) -> None:
        template_path = self.temp_dir / "template.docx"
        output_text_path = self.temp_dir / "result.txt"
        output_docx_path = self.temp_dir / "result.docx"
        self.write_template_docx(template_path)
        output_text_path.write_text("只有一段。", encoding="utf-8")

        with self.assertRaisesRegex(ValueError, "paragraph count"):
            render_docx_from_template(template_path, output_text_path, output_docx_path)

        self.assertFalse(output_docx_path.exists())

    def test_render_docx_from_template_does_not_replace_existing_target_on_failure(self) -> None:
        template_path = self.temp_dir / "template.docx"
        output_text_path = self.temp_dir / "result.txt"
        output_docx_path = self.temp_dir / "result.docx"
        self.write_template_docx(template_path)
        output_text_path.write_text("只有一段。", encoding="utf-8")
        output_docx_path.write_bytes(b"existing output")

        with self.assertRaisesRegex(ValueError, "paragraph count"):
            render_docx_from_template(template_path, output_text_path, output_docx_path)

        self.assertEqual(output_docx_path.read_bytes(), b"existing output")

    def test_skill_round_materializes_preserved_docx_and_records_its_path(self) -> None:
        source_path = self.temp_dir / "source.docx"
        self.write_template_docx(source_path)
        finish_dir = self.temp_dir / "finish"
        intermediate_dir = finish_dir / "intermediate"
        records_path = finish_dir / "aigc_records.json"

        with (
            patch("skill_round_helper.INTERMEDIATE_DIR", intermediate_dir),
            patch("aigc_records.FINISH_DIR", finish_dir),
            patch("aigc_records.RECORDS_PATH", records_path),
        ):
            result = run_skill_round(
                source_path,
                transform=lambda chunk_text, _prompt, _round, _chunk_id: chunk_text,
            )

        docx_output_path = Path(result["skill_context"]["docx_output_path"])
        self.assertTrue(docx_output_path.exists())
        self.assertEqual(
            [paragraph.text for paragraph in Document(str(docx_output_path)).paragraphs],
            ["原始第一段。", "原始第二段。"],
        )

        records = json.loads(records_path.read_text(encoding="utf-8"))
        record_key = str(source_path.resolve()).replace("\\", "/")
        round_record = records[record_key]["rounds"][0]
        self.assertEqual(round_record["docx_output_path"], docx_output_path.as_posix())

    def test_app_export_uses_preserved_docx_artifact(self) -> None:
        source_path = self.temp_dir / "source.docx"
        output_text_path = self.temp_dir / "result.txt"
        generated_docx_path = self.temp_dir / "generated.docx"
        export_path = self.temp_dir / "exported.docx"
        self.write_template_docx(source_path)
        output_text_path.write_text("导出的第一段。\n\n导出的第二段。", encoding="utf-8")

        render_docx_from_template(source_path, output_text_path, generated_docx_path)
        result = export_round_output(
            str(output_text_path),
            str(export_path),
            "docx",
            source_path=str(source_path),
            docx_output_path=str(generated_docx_path),
        )

        self.assertEqual(result["path"], str(export_path.resolve()))
        document = Document(str(export_path))
        self.assertEqual(document.paragraphs[0].text, "导出的第一段。")
        self.assertTrue(document.paragraphs[0].runs[0].bold)
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(document.sections[0].header.paragraphs[0].text, "页眉保留")

    def test_app_export_lazily_renders_docx_from_original_template(self) -> None:
        source_path = self.temp_dir / "source.docx"
        output_text_path = self.temp_dir / "result.txt"
        manifest_path = self.temp_dir / "result_manifest.json"
        export_path = self.temp_dir / "exported.docx"
        self.write_template_docx(source_path)
        output_text_path.write_text("延迟生成的第一段。\n\n延迟生成的第二段。", encoding="utf-8")
        manifest_path.write_text(json.dumps({"paragraph_count": 2}), encoding="utf-8")

        result = export_round_output(
            str(output_text_path),
            str(export_path),
            "docx",
            source_path=str(source_path),
            manifest_path=str(manifest_path),
        )

        self.assertEqual(result["path"], str(export_path.resolve()))
        document = Document(str(export_path))
        self.assertEqual(
            [paragraph.text for paragraph in document.paragraphs],
            ["延迟生成的第一段。", "延迟生成的第二段。"],
        )
        self.assertTrue(document.paragraphs[0].runs[0].bold)
        self.assertEqual(document.tables[0].cell(0, 0).text, "表格内容保持不变")
        self.assertEqual(document.sections[0].header.paragraphs[0].text, "页眉保留")

    def test_txt_source_keeps_plain_docx_export_behavior(self) -> None:
        source_path = self.temp_dir / "source.txt"
        output_text_path = self.temp_dir / "result.txt"
        export_path = self.temp_dir / "exported.docx"
        source_path.write_text("原始文本。", encoding="utf-8")
        output_text_path.write_text("TXT 导出的内容。", encoding="utf-8")

        result = export_round_output(
            str(output_text_path),
            str(export_path),
            "docx",
            source_path=str(source_path),
        )

        self.assertEqual(result["format"], "docx")
        self.assertEqual([paragraph.text for paragraph in Document(str(export_path)).paragraphs], ["TXT 导出的内容。"])

    def test_app_export_rejects_docx_without_source_template(self) -> None:
        output_text_path = self.temp_dir / "result.txt"
        export_path = self.temp_dir / "exported.docx"
        output_text_path.write_text("导出内容。", encoding="utf-8")

        with self.assertRaisesRegex(ValueError, "original DOCX source path"):
            export_round_output(str(output_text_path), str(export_path), "docx")

        self.assertFalse(export_path.exists())

    def test_app_round_returns_preserved_docx_output_for_docx_source(self) -> None:
        source_path = self.temp_dir / "source.docx"
        finish_dir = self.temp_dir / "finish"
        intermediate_dir = finish_dir / "intermediate"
        records_path = finish_dir / "aigc_records.json"
        self.write_template_docx(source_path)

        with (
            patch("skill_round_helper.INTERMEDIATE_DIR", intermediate_dir),
            patch("aigc_records.FINISH_DIR", finish_dir),
            patch("aigc_records.RECORDS_PATH", records_path),
        ):
            result = run_round_for_app(
                str(source_path),
                {
                    "offlineMode": True,
                    "promptProfile": "cn",
                },
            )

        docx_output_path = Path(result["docxOutputPath"])
        self.assertTrue(docx_output_path.exists())
        self.assertEqual(
            [paragraph.text for paragraph in Document(str(docx_output_path)).paragraphs],
            ["原始第一段。", "原始第二段。"],
        )
        self.assertEqual(result["skillContext"]["docx_output_path"], str(docx_output_path))

        records = json.loads(records_path.read_text(encoding="utf-8"))
        record_key = str(source_path.resolve()).replace("\\", "/")
        self.assertEqual(
            records[record_key]["rounds"][0]["docx_output_path"],
            docx_output_path.as_posix(),
        )


if __name__ == "__main__":
    unittest.main()
