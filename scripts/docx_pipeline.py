"""Utility script for reading and writing .docx files
for the baibaiaigc skill.

This script does NOT run the three-step AIGC reduction itself.
It only helps you move text between .docx files and plain text,
so the skill can work on the text while you keep a .docx workflow.

Requirements:
  pip install python-docx

Typical usage:
  1. Put an input .docx file under the workspace root `origin/` directory.
  2. Run this script to extract its plain text.
  3. Use the baibaiaigc skill on the extracted text.
  4. Optionally, write the final text back into a new .docx file
     under the `finish/` directory.
"""

from __future__ import annotations

import argparse
import json
import os
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Iterable

try:
    from docx import Document  # type: ignore[import]
except ImportError as exc:  # pragma: no cover - import guard
    raise SystemExit(
        "Missing dependency python-docx. Install it with: pip install python-docx"
    ) from exc


WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
XML_NAMESPACE = "http://www.w3.org/XML/1998/namespace"
MAIN_DOCUMENT_PART = "word/document.xml"
W_BODY = f"{{{WORD_NAMESPACE}}}body"
W_PARAGRAPH = f"{{{WORD_NAMESPACE}}}p"
W_TEXT = f"{{{WORD_NAMESPACE}}}t"
XML_SPACE = f"{{{XML_NAMESPACE}}}space"


def read_docx_text(path: Path) -> str:
    """Read a .docx file and return its text as paragraphs joined by blank lines."""
    document = Document(str(path))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    # Keep empty lines only when there is actual content separation.
    non_empty_blocks: list[str] = []
    for paragraph in paragraphs:
        if paragraph:
            non_empty_blocks.append(paragraph)
    return "\n\n".join(non_empty_blocks)


def _main_body_paragraphs(root: ET.Element) -> list[ET.Element]:
    body = root.find(f".//{W_BODY}")
    if body is None:
        raise ValueError("DOCX main document body is missing.")
    return [child for child in list(body) if child.tag == W_PARAGRAPH]


def _paragraph_xml_text(paragraph: ET.Element) -> str:
    return "".join(node.text or "" for node in paragraph.iter(W_TEXT))


def _replace_paragraph_xml_text(paragraph: ET.Element, new_text: str) -> None:
    text_nodes = list(paragraph.iter(W_TEXT))
    if not text_nodes:
        raise ValueError("DOCX paragraph has no writable text node.")

    old_lengths = [len(node.text or "") for node in text_nodes]
    old_total = sum(old_lengths)
    remaining = new_text
    for index, node in enumerate(text_nodes):
        if index == len(text_nodes) - 1:
            value = remaining
        elif old_total:
            target_length = round(len(new_text) * old_lengths[index] / old_total)
            value = remaining[:target_length]
            remaining = remaining[target_length:]
        else:
            value = ""

        node.text = value
        if value.startswith((" ", "\t")) or value.endswith((" ", "\t")):
            node.set(XML_SPACE, "preserve")
        else:
            node.attrib.pop(XML_SPACE, None)


def _read_manifest_paragraph_count(path: Path) -> int:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"Invalid DOCX manifest: {path}") from exc
    if not isinstance(data, dict) or not isinstance(data.get("paragraph_count"), int):
        raise ValueError(f"DOCX manifest has no paragraph count: {path}")
    return int(data["paragraph_count"])


def _validate_docx_paragraph_counts(
    template_count: int,
    output_count: int,
    manifest_path: Path | None,
) -> None:
    if template_count != output_count:
        raise ValueError(
            f"DOCX paragraph count mismatch: template has {template_count}, "
            f"output has {output_count}."
        )
    if manifest_path is not None:
        manifest_count = _read_manifest_paragraph_count(manifest_path)
        if manifest_count != template_count:
            raise ValueError(
                f"DOCX paragraph count mismatch: manifest has {manifest_count}, "
                f"template has {template_count}."
            )


def render_docx_from_template(
    template_path: Path,
    output_text_path: Path,
    output_docx_path: Path,
    manifest_path: Path | None = None,
) -> Path:
    """Render text into an existing DOCX package without rebuilding it."""
    normalized_template_path = Path(template_path).resolve()
    normalized_output_text_path = Path(output_text_path).resolve()
    normalized_output_docx_path = Path(output_docx_path).resolve()
    normalized_manifest_path = Path(manifest_path).resolve() if manifest_path is not None else None

    if normalized_template_path.suffix.lower() != ".docx":
        raise ValueError(f"DOCX template must be a .docx file: {normalized_template_path}")
    if not normalized_template_path.is_file():
        raise ValueError(f"DOCX template not found: {normalized_template_path}")
    if not normalized_output_text_path.is_file():
        raise ValueError(f"DOCX output text not found: {normalized_output_text_path}")
    if normalized_manifest_path is not None and not normalized_manifest_path.is_file():
        raise ValueError(f"DOCX manifest not found: {normalized_manifest_path}")
    if normalized_template_path == normalized_output_docx_path:
        raise ValueError("DOCX output must not overwrite the template.")

    output_blocks = _split_text_into_blocks(
        normalized_output_text_path.read_text(encoding="utf-8")
    )

    ET.register_namespace("w", WORD_NAMESPACE)
    temp_path: Path | None = None
    try:
        with zipfile.ZipFile(normalized_template_path, "r") as source_archive:
            if MAIN_DOCUMENT_PART not in source_archive.namelist():
                raise ValueError("DOCX main document part is missing.")

            root = ET.fromstring(source_archive.read(MAIN_DOCUMENT_PART))
            source_paragraphs = [
                paragraph
                for paragraph in _main_body_paragraphs(root)
                if _paragraph_xml_text(paragraph).strip()
            ]
            _validate_docx_paragraph_counts(
                len(source_paragraphs),
                len(output_blocks),
                normalized_manifest_path,
            )
            for paragraph, output_text in zip(source_paragraphs, output_blocks):
                _replace_paragraph_xml_text(paragraph, output_text)
            updated_document = ET.tostring(root, encoding="utf-8", xml_declaration=True)

            normalized_output_docx_path.parent.mkdir(parents=True, exist_ok=True)
            file_descriptor, temporary_name = tempfile.mkstemp(
                prefix=f"{normalized_output_docx_path.stem}-",
                suffix=".docx",
                dir=normalized_output_docx_path.parent,
            )
            os.close(file_descriptor)
            temp_path = Path(temporary_name)

            with zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target_archive:
                for item in source_archive.infolist():
                    data = (
                        updated_document
                        if item.filename == MAIN_DOCUMENT_PART
                        else source_archive.read(item.filename)
                    )
                    target_archive.writestr(item, data)

        os.replace(temp_path, normalized_output_docx_path)
        temp_path = None
        return normalized_output_docx_path
    finally:
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)


def read_docx_paragraphs(path: Path) -> list[str]:
    """Read a .docx file and return non-empty paragraph texts in order."""
    document = Document(str(path))
    paragraphs: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            paragraphs.append(text)
    return paragraphs


def write_docx_text(lines: Iterable[str], path: Path) -> None:
    """Write an iterable of text blocks into a .docx file.

    Each element in `lines` becomes one paragraph.
    """
    document = Document()
    for block in lines:
        document.add_paragraph(block)
    document.save(str(path))


def write_docx_paragraphs(paragraphs: Iterable[str], path: Path) -> None:
    """Write one paragraph per entry, preserving paragraph boundaries."""
    write_docx_text(paragraphs, path)


def _split_text_into_blocks(text: str) -> list[str]:
    """Split a large text into paragraph blocks using blank lines as separators."""
    blocks: list[str] = []
    current: list[str] = []
    for raw_line in text.splitlines():
        line = raw_line.rstrip("\n")
        if not line.strip():
            if current:
                blocks.append(" ".join(current).strip())
                current = []
            continue
        current.append(line.strip())
    if current:
        blocks.append(" ".join(current).strip())
    return blocks


def _read_paragraphs_file(path: Path) -> list[str]:
    if path.suffix.lower() == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, list) or not all(isinstance(item, str) for item in data):
            raise SystemExit("Paragraph json must be a string array.")
        return data
    text = path.read_text(encoding="utf-8")
    return _split_text_into_blocks(text)


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description="Simple .docx <-> text helper")
    subparsers = parser.add_subparsers(dest="command", required=True)

    extract_parser = subparsers.add_parser(
        "extract", help="Extract plain text from a .docx file",
    )
    extract_parser.add_argument("input", type=Path, help="Path to input .docx file")

    extract_to_file_parser = subparsers.add_parser(
        "extract-to-file", help="Extract plain text from a .docx file into a UTF-8 text file",
    )
    extract_to_file_parser.add_argument(
        "input", type=Path, help="Path to input .docx file"
    )
    extract_to_file_parser.add_argument(
        "output", type=Path, help="Path to output .txt file"
    )

    extract_paragraphs_parser = subparsers.add_parser(
        "extract-paragraphs",
        help="Extract non-empty paragraphs from a .docx file into a JSON array",
    )
    extract_paragraphs_parser.add_argument(
        "input", type=Path, help="Path to input .docx file"
    )
    extract_paragraphs_parser.add_argument(
        "output", type=Path, help="Path to output .json file"
    )

    build_parser = subparsers.add_parser(
        "build", help="Build a .docx file from a plain text file",
    )
    build_parser.add_argument("input", type=Path, help="Path to input .txt file")
    build_parser.add_argument("output", type=Path, help="Path to output .docx file")

    build_paragraphs_parser = subparsers.add_parser(
        "build-paragraphs",
        help="Build a .docx file from a paragraph JSON array or block text file",
    )
    build_paragraphs_parser.add_argument("input", type=Path, help="Path to paragraph json/txt file")
    build_paragraphs_parser.add_argument("output", type=Path, help="Path to output .docx file")

    args = parser.parse_args(argv)

    if args.command == "extract":
        text = read_docx_text(args.input)
        print(text)
    elif args.command == "extract-to-file":
        text = read_docx_text(args.input)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(text, encoding="utf-8")
    elif args.command == "extract-paragraphs":
        paragraphs = read_docx_paragraphs(args.input)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(
            json.dumps(paragraphs, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    elif args.command == "build":
        text = args.input.read_text(encoding="utf-8")
        blocks = _split_text_into_blocks(text)
        write_docx_text(blocks, args.output)
    elif args.command == "build-paragraphs":
        paragraphs = _read_paragraphs_file(args.input)
        write_docx_paragraphs(paragraphs, args.output)
    else:  # pragma: no cover - argparse guarantees command
        parser.error("Unknown command")


if __name__ == "__main__":  # pragma: no cover - CLI entry point
    main()
